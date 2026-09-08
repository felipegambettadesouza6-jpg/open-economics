#!/usr/bin/env python3
"""Build a versioned Government Central fiscal-series snapshot from the official RTN workbook."""

from __future__ import annotations

import argparse
import json
import re
import tempfile
import urllib.request
from datetime import date, datetime, timezone
from pathlib import Path

import openpyxl


PACKAGE_API = "https://www.tesourotransparente.gov.br/ckan/api/3/action/package_show?id=resultado-do-tesouro-nacional"
SOURCE_PAGE = "https://www.tesourotransparente.gov.br/ckan/dataset/resultado-do-tesouro-nacional"
OUTPUT = Path(__file__).resolve().parents[1] / "lib" / "catalog" / "generated" / "tesouro-rtn.generated.json"
USER_AGENT = "Mozilla/5.0 OpenEconomicsCatalogSync/2.0"


def fetch_json(url: str) -> dict:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=60) as response:
        return json.load(response)


def download(url: str, target: Path) -> None:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=120) as response, target.open("wb") as output:
        while chunk := response.read(1024 * 1024):
            output.write(chunk)


def account_id(label: str) -> str | None:
    match = re.match(r"^\s*(\d+(?:\.\d+)*)\.?\s", label)
    return match.group(1) if match else None


def numeric(value: object) -> int | float | None:
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return None
    parsed = float(value)
    return int(parsed) if parsed.is_integer() else round(parsed, 8)


def parse_workbook(path: Path) -> tuple[list[list[object]], list[list[object]], list[str], str, str]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "1.2" not in workbook.sheetnames:
            raise RuntimeError("RTN workbook no longer contains detailed current-value table 1.2.")
        sheet = workbook["1.2"]
        matrix = list(sheet.iter_rows(values_only=True))
        title = str(matrix[1][0] or "")
        unit = str(matrix[2][0] or "")
        if "Resultado Primário do Governo Central" not in title or "Valores Correntes" not in unit:
            raise RuntimeError("RTN table 1.2 title or unit changed unexpectedly.")
        dates: list[tuple[int, str]] = []
        for column, value in enumerate(matrix[4][1:], 1):
            if isinstance(value, (datetime, date)):
                dates.append((column, value.strftime("%Y%m")))
        if not dates or dates[0][1] != "199701":
            raise RuntimeError("RTN monthly coverage no longer starts in January 1997.")

        accounts: list[list[object]] = []
        observations: list[list[object]] = []
        seen: set[str] = set()
        for row in matrix[5:164]:
            label = str(row[0] or "").strip()
            identifier = account_id(label)
            if not identifier:
                continue
            if identifier in seen:
                raise RuntimeError(f"Duplicate RTN account identifier: {identifier}")
            seen.add(identifier)
            accounts.append([identifier, label, identifier.count(".") + 1, identifier.split(".")[0]])
            for column, period in dates:
                raw = row[column]
                value = numeric(raw)
                status = "observed" if value is not None else "not-available" if isinstance(raw, str) and raw.strip() else "missing"
                observations.append([period, identifier, value, status])

        footnotes = [str(row[0]).strip() for row in matrix[178:] if row[0]]
    finally:
        workbook.close()
    if len(accounts) < 150 or not observations:
        raise RuntimeError("RTN account coverage changed unexpectedly.")
    return accounts, observations, footnotes, title, unit


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, help="Optional already-downloaded official RTN workbook.")
    arguments = parser.parse_args()
    package = fetch_json(PACKAGE_API)["result"]
    data_resource = next(resource for resource in package["resources"] if resource.get("name") == "Resultado do Tesouro Nacional - Série Histórica - Mensal")
    dictionary_resource = next(resource for resource in package["resources"] if resource.get("name") == "Dicionário de Conceitos e Metodologia de Cálculo")
    metadata_resource = next(resource for resource in package["resources"] if resource.get("name") == "Metadados")

    if arguments.workbook:
        path = arguments.workbook.resolve()
        if not path.is_file():
            raise RuntimeError("--workbook must point to an official RTN XLSX file.")
        accounts, observations, footnotes, title, unit = parse_workbook(path)
    else:
        with tempfile.TemporaryDirectory(prefix="open-economics-rtn-") as directory:
            path = Path(directory) / "rtn.xlsx"
            download(data_resource["url"], path)
            accounts, observations, footnotes, title, unit = parse_workbook(path)

    periods = [row[0] for row in observations]
    payload = {
        "schema_version": 1,
        "synced_at": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "source_version": package["metadata_modified"],
        "source_file": data_resource["url"].rsplit("/", 1)[-1],
        "source_url": data_resource["url"],
        "source_page": SOURCE_PAGE,
        "package_api": PACKAGE_API,
        "dictionary_url": dictionary_resource["url"],
        "metadata_url": metadata_resource["url"],
        "table": "1.2",
        "title": title,
        "unit": unit,
        "coverage": {"start": min(periods), "end": max(periods)},
        "account_columns": ["id", "label", "depth", "section"],
        "accounts": accounts,
        "observation_columns": ["period", "account_id", "value_millions_brl", "status"],
        "observations": observations,
        "footnotes": footnotes,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT), "accounts": len(accounts), "observations": len(observations), "coverage": payload["coverage"], "bytes": OUTPUT.stat().st_size}, ensure_ascii=False))


if __name__ == "__main__":
    main()
