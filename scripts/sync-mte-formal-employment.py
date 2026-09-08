#!/usr/bin/env python3
"""Build a compact, versioned Novo Caged snapshot from the latest official table."""

from __future__ import annotations

import html
import argparse
import http.cookiejar
import json
import re
import tempfile
import unicodedata
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path

import openpyxl


ROOT_FOLDER_ID = "1F89h6odTPGIGMb9eDiJKCute9W89QmqN"
SOURCE_PAGE = "https://www.gov.br/trabalho-e-emprego/pt-br/assuntos/estatisticas-trabalho/novo-caged"
FOLDER_URL = f"https://drive.google.com/drive/folders/{ROOT_FOLDER_ID}"
OUTPUT = Path(__file__).resolve().parents[1] / "lib" / "catalog" / "generated" / "mte-formal-employment.generated.json"
USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/128.0 Safari/537.36 OpenEconomicsCatalogSync/2.0"

MONTHS = {
    "janeiro": 1, "fevereiro": 2, "marco": 3, "abril": 4, "maio": 5, "junho": 6,
    "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12,
}
REGIONS = {"Norte", "Nordeste", "Sudeste", "Sul", "Centro-Oeste"}
STATE_CODES = {
    "Acre": "AC", "Alagoas": "AL", "Amapá": "AP", "Amazonas": "AM", "Bahia": "BA",
    "Ceará": "CE", "Distrito Federal": "DF", "Espírito Santo": "ES", "Goiás": "GO",
    "Maranhão": "MA", "Mato Grosso": "MT", "Mato Grosso do Sul": "MS", "Minas Gerais": "MG",
    "Pará": "PA", "Paraíba": "PB", "Paraná": "PR", "Pernambuco": "PE", "Piauí": "PI",
    "Rio de Janeiro": "RJ", "Rio Grande do Norte": "RN", "Rio Grande do Sul": "RS",
    "Rondônia": "RO", "Roraima": "RR", "Santa Catarina": "SC", "São Paulo": "SP",
    "Sergipe": "SE", "Tocantins": "TO",
}


def normalized(value: object) -> str:
    return "".join(
        character for character in unicodedata.normalize("NFKD", str(value or "").lower())
        if not unicodedata.combining(character)
    ).strip()


def period(value: object) -> str | None:
    match = re.fullmatch(r"\s*([^/]+)/([0-9]{4})\s*", str(value or ""))
    if not match:
        return None
    month = MONTHS.get(normalized(match.group(1)))
    return f"{match.group(2)}{month:02d}" if month else None


def number(value: object) -> int | float | None:
    if value is None or isinstance(value, bool) or str(value).strip() in {"", "-", "..."}:
        return None
    parsed = float(value)
    return int(parsed) if parsed.is_integer() else round(parsed, 8)


def slug(value: str) -> str:
    return re.sub(r"(^-|-$)", "", re.sub(r"[^a-z0-9]+", "-", normalized(value)))


class LinkParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.href: str | None = None
        self.text: list[str] = []
        self.links: list[tuple[str, str]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "a":
            self.href = dict(attrs).get("href")
            self.text = []

    def handle_data(self, data: str) -> None:
        if self.href:
            self.text.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag == "a" and self.href:
            self.links.append((self.href, html.unescape("".join(self.text).strip())))
            self.href = None
            self.text = []


def fetch_text(url: str) -> str:
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=30) as response:
        return response.read().decode("utf-8")


def list_folder(folder_id: str, prefix: str = "") -> list[tuple[str, str]]:
    url = "https://drive.google.com/embeddedfolderview?" + urllib.parse.urlencode({"id": folder_id})
    parser = LinkParser()
    parser.feed(fetch_text(url))
    files: list[tuple[str, str]] = []
    for href, name in parser.links:
        folder_match = re.match(r"https://drive\.google\.com/drive/folders/([-\w]{25,})", href)
        file_match = re.match(r"https://drive\.google\.com/file/d/([-\w]{25,})/view", href)
        if folder_match:
            files.extend(list_folder(folder_match.group(1), f"{prefix}{name}/"))
        elif file_match:
            files.append((file_match.group(1), f"{prefix}{name}"))
    return files


def latest_workbook() -> tuple[str, str, str]:
    candidates: list[tuple[str, str, str]] = []
    for file_id, path in list_folder(ROOT_FOLDER_ID):
        match = re.search(r"(?:^|/)(20\d{4})(?:/|$).+\.xlsx$", path, re.IGNORECASE)
        if match:
            candidates.append((match.group(1), file_id, path))
    if not candidates:
        raise RuntimeError("The official Novo Caged folder did not contain a versioned XLSX workbook.")
    return max(candidates)


def download(file_id: str, target: Path) -> str:
    url = "https://drive.google.com/uc?" + urllib.parse.urlencode({"export": "download", "id": file_id, "confirm": "t"})
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
    next_url = url
    for _ in range(3):
        request = urllib.request.Request(next_url, headers={"User-Agent": USER_AGENT})
        with opener.open(request, timeout=120) as response:
            prefix = response.read(4)
            if prefix == b"PK\x03\x04":
                with target.open("wb") as output:
                    output.write(prefix)
                    while chunk := response.read(1024 * 1024):
                        output.write(chunk)
                return url
            contents = prefix + response.read()
        page = contents.decode("utf-8", errors="replace")
        form = re.search(r'<form[^>]+id="download-form"[^>]+action="([^"]+)"[\s\S]*?</form>', page)
        if not form:
            break
        fields = dict(re.findall(r'<input[^>]+type="hidden"[^>]+name="([^"]+)"[^>]+value="([^"]*)"', form.group(0)))
        next_url = html.unescape(form.group(1)) + "?" + urllib.parse.urlencode({key: html.unescape(value) for key, value in fields.items()})
    raise RuntimeError("Google Drive did not return an XLSX workbook. Retry later or pass --workbook with the downloaded official file.")


def monthly_columns(matrix: list[tuple[object, ...]], title: str) -> list[tuple[str, dict[str, int]]]:
    header = matrix[4]
    measures_row = matrix[5]
    boundaries = [column for column in range(1, len(header)) if header[column] is not None]
    starts = [(column, period(header[column])) for column in range(1, len(header))]
    starts = [(column, value) for column, value in starts if value]
    result: list[tuple[str, dict[str, int]]] = []
    for start, value in starts:
        end = next((column for column in boundaries if column > start), len(measures_row))
        measures: dict[str, int] = {}
        for column in range(start, end):
            label = normalized(measures_row[column])
            if label.startswith("estoque"):
                measures["stock"] = column
            elif label.startswith("admisso"):
                measures["admissions"] = column
            elif label.startswith("desligamento"):
                measures["dismissals"] = column
            elif label.startswith("saldo"):
                measures["balance"] = column
            elif label.startswith("variacao relativa"):
                measures["relative_change_pct"] = column
        if not {"stock", "admissions", "dismissals", "balance"}.issubset(measures):
            raise RuntimeError(f"Unexpected Novo Caged measure columns for {value} in {title}.")
        result.append((value, measures))
    return result


def record(value: str, scope: str, entity_id: str, label: str, row: tuple[object, ...], measures: dict[str, int]) -> list[object]:
    return [
        value, scope, entity_id, label,
        number(row[measures["stock"]]),
        number(row[measures["admissions"]]),
        number(row[measures["dismissals"]]),
        number(row[measures["balance"]]),
        number(row[measures["relative_change_pct"]]) if "relative_change_pct" in measures else None,
    ]


def parse_workbook(path: Path) -> list[list[object]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        required = {"Tabela 5.1", "Tabela 6.1", "Tabela 7.1"}
        if not required.issubset(workbook.sheetnames):
            raise RuntimeError("Novo Caged workbook is missing an adjusted historical table.")
        rows: list[list[object]] = []

        national = list(workbook["Tabela 5.1"].iter_rows(values_only=True))
        if "com ajustes" not in normalized(national[1][1]):
            raise RuntimeError("Tabela 5.1 no longer identifies the adjusted series.")
        for row in national[5:]:
            value = period(row[1])
            if not value:
                continue
            measures = {"stock": 2, "admissions": 3, "dismissals": 4, "balance": 5, "relative_change_pct": 6}
            rows.append(record(value, "country", "BR", "Brasil", row, measures))

        industry = list(workbook["Tabela 6.1"].iter_rows(values_only=True))
        industry_months = monthly_columns(industry, "Tabela 6.1")
        for row in industry[6:]:
            label = str(row[1] or "").strip()
            if not label or normalized(label) in {"total", "nao identificado***", "nao identificado"} or label.startswith(("Fonte:", "*")):
                continue
            for value, measures in industry_months:
                rows.append(record(value, "industry", slug(label), label, row, measures))

        geography = list(workbook["Tabela 7.1"].iter_rows(values_only=True))
        geography_months = monthly_columns(geography, "Tabela 7.1")
        for row in geography[6:]:
            label = str(row[1] or "").strip()
            if not label or label in {"Brasil", "Não identificado"} or label.startswith(("Fonte:", "*")):
                continue
            if label in REGIONS:
                scope, entity_id = "region", slug(label)
            elif label in STATE_CODES:
                scope, entity_id = "state", STATE_CODES[label]
            else:
                continue
            for value, measures in geography_months:
                rows.append(record(value, scope, entity_id, label, row, measures))
    finally:
        workbook.close()

    rows.sort(key=lambda item: (item[0], {"country": 0, "region": 1, "state": 2, "industry": 3}[item[1]], item[2]))
    periods = sorted({row[0] for row in rows})
    if not periods or periods[0] != "202001" or len([row for row in rows if row[1] == "state"]) != len(periods) * 27:
        raise RuntimeError("Novo Caged adjusted-series coverage changed unexpectedly.")
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, help="Optional already-downloaded official XLSX; latest folder metadata is still verified.")
    arguments = parser.parse_args()
    source_version, file_id, source_path = latest_workbook()
    workbook_url = "https://drive.google.com/uc?" + urllib.parse.urlencode({"export": "download", "id": file_id, "confirm": "t"})
    if arguments.workbook:
        workbook_path = arguments.workbook.resolve()
        if not workbook_path.is_file():
            raise RuntimeError("--workbook must point to a valid official XLSX file.")
        with workbook_path.open("rb") as workbook_stream:
            if workbook_stream.read(4) != b"PK\x03\x04":
                raise RuntimeError("--workbook must point to a valid official XLSX file.")
        rows = parse_workbook(workbook_path)
    else:
        with tempfile.TemporaryDirectory(prefix="open-economics-mte-") as directory:
            workbook_path = Path(directory) / "novo-caged.xlsx"
            workbook_url = download(file_id, workbook_path)
            rows = parse_workbook(workbook_path)
    payload = {
        "schema_version": 1,
        "synced_at": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
        "source_version": source_version,
        "source_file": source_path,
        "source_url": workbook_url,
        "source_page": SOURCE_PAGE,
        "folder_url": FOLDER_URL,
        "columns": ["period", "scope", "id", "label", "stock", "admissions", "dismissals", "balance", "relative_change_pct"],
        "rows": rows,
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT), "rows": len(rows), "source_version": source_version, "source_file": source_path}, ensure_ascii=False))


if __name__ == "__main__":
    main()
